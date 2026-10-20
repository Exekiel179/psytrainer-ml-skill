#!/usr/bin/env python3
"""Real Pipeline training, prediction and Word output on synthetic test data."""

import argparse
import json
import subprocess
import sys
import tempfile
import zipfile
from pathlib import Path

import numpy as np
import pandas as pd
from docx import Document

ROOT = Path(__file__).resolve().parents[1]


def exercise(root):
    root.mkdir(parents=True, exist_ok=True)
    rng = np.random.default_rng(13)
    ids = [f"{i:04}" for i in range(120)]
    x = pd.DataFrame(rng.normal(size=(120, 4)), index=ids, columns=["signal", "noise", "age", "baseline"])
    y = pd.DataFrame({"score": 2 * x.signal - x.baseline + rng.normal(size=120) * .2,
                      "class": np.where(x.signal > 0, "positive", "negative")}, index=ids)
    x.iloc[::13, 1] = np.nan
    m = pd.DataFrame({"subject": np.repeat(np.arange(30), 4), "time": np.repeat(np.arange(60), 2)}, index=ids)
    for name, frame in (("x", x), ("y", y), ("meta", m)):
        frame.to_csv(root / f"{name}.csv")

    def command(*args):
        completed = subprocess.run([sys.executable, str(ROOT / "scripts/pipeline_train.py"), *map(str, args)],
                                   capture_output=True, text=True)
        if completed.returncode:
            raise RuntimeError(completed.stdout + completed.stderr)
        return json.loads(completed.stdout)

    for split, task, target in (("random", "regression", "score"), ("group", "classification", "class"),
                                 ("time", "regression", "score")):
        out = root / split
        extra = [] if split == "random" else ["--metadata", root / "meta.csv", "--group-column" if split == "group" else "--time-column",
                                               "subject" if split == "group" else "time"]
        if split == "group":
            extra += ["--resample", "smote", "--search", "random", "--max-candidates", 2,
                      "--metric", "precision"]
        if split == "time":
            extra += ["--gap", 1, "--selection", "forward", "--metric", "pearson_r"]
        if split == "random":
            extra += ["--scaler", "minmax", "--search", "grid", "--metric", "mse"]
        result = command("train", "--features", root / "x.csv", "--labels", root / "y.csv", "--target", target,
                         "--task", task, "--split", split, "--output-dir", out, "--cv", 3, "--repeats", 3,
                         "--impute", "median", "--select-k", 3, "--pca", 2, *extra)
        assert result["status"] == "completed" and not result["failures"], result
        summary = json.loads((out / "summary.json").read_text(encoding="utf-8"))
        assert summary["development_n"] + summary["test_n"] <= len(x)
        assert len(summary["comparison"]) == 2
        if split != "time":
            assert summary["search"]["method"] != "none"
            assert list((out / "figures").glob("parameter-search-*.svg"))
        else:
            assert summary["preprocessing"]["selection"] == "forward"
        diagnostics = json.loads((out / "analysis.json").read_text(encoding="utf-8"))
        assert len(diagnostics["findings"]) >= 4
        intervals = pd.read_csv(out / "metric-intervals.csv")
        if split == "time":
            assert intervals.lower.isna().all()
        else:
            assert intervals.lower.notna().any()
        oof = pd.read_csv(out / "oof-predictions.csv", dtype={"sample_id": str})
        membership = pd.read_csv(out / "split-membership.csv", dtype={"sample_id": str})
        test_ids = set(membership.loc[membership.partition == "test", "sample_id"])
        assert not set(oof.sample_id) & test_ids
        assert len(pd.read_csv(out / "error-cases.csv")) == summary["test_n"]
        with zipfile.ZipFile(out / "report.docx") as archive:
            assert archive.testzip() is None
        doc = Document(out / "report.docx")
        assert len(doc.inline_shapes) >= 4 and doc.tables
        assert len(list((out / "figures").glob("*.svg"))) == len(doc.inline_shapes)
        headings = [p.text for p in doc.paragraphs if p.style.name.startswith("Heading")]
        assert any(h in ("2.1 模型比较", "2.1 Model comparison") for h in headings)
        assert "摘要" in headings or "Abstract" in headings
        md = (out / "report.md").read_text(encoding="utf-8")
        assert "附录 A" in md or "Appendix A" in md
        assert "主张：" not in md and "Claim:" not in md
        sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "scripts"))
        from report_narrative import fmt, metric_label
        zh = summary["language"] == "zh"
        test_table = doc.tables[1]
        for metric, value in summary["test_metrics"].items():
            if value is not None:
                assert any(row.cells[0].text.startswith(metric_label(metric, zh)) and row.cells[1].text == fmt(value)
                           for row in test_table.rows), metric
        command("predict", "--features", root / "x.csv", "--model", out / "pipeline.joblib", "--output", out / "new-predictions.csv")
        assert len(pd.read_csv(out / "new-predictions.csv")) == len(x)
        print(json.dumps({"case": split, "status": "passed", "report": result["report"]}))


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--output", type=Path)
    args = parser.parse_args()
    if args.output:
        exercise(args.output.resolve())
    else:
        with tempfile.TemporaryDirectory(prefix="psytrainer-pipeline-") as temp:
            exercise(Path(temp))
