"""Structure checks for academic/technical Word and Markdown reports."""

from __future__ import annotations

import json
import sys
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

import numpy as np
import pandas as pd

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "scripts"))
import pipeline_report
from report_explanations import PER_MODEL_SPLIT_THRESHOLD


class PipelineReportTests(unittest.TestCase):
    def _write_run(self, root: Path, models: list[str], *, selected: str | None = None):
        selected = selected or models[0]
        comparison = [{"model": tag, "mean": 0.2 + i * 0.05, "sd": 0.01, "folds": 3}
                      for i, tag in enumerate(models)]
        summary = {
            "schema": "psytrainer-pipeline/v1", "task": "regression", "target": "score",
            "split": "random", "external_test": False, "seed": 42, "cv_folds": 3, "gap": 0,
            "development_n": 60, "test_n": 20, "feature_n": 3, "metric": "mae",
            "direction": "lower", "selected_model": selected, "comparison": comparison,
            "test_metrics": {"mae": 0.25, "mse": 0.1, "rmse": 0.3162, "r2": 0.8, "pearson_r": 0.9},
            "baseline": {"strategy": "mean", "test_metrics": {"mae": 0.5}},
            "bootstrap_repeats": 100, "score_kind": "none", "classes": [], "class_labels": None,
            "search": {"method": "none", "max_candidates": 0, "evaluated_candidates": 0,
                       "selected_parameters": {}, "failed_candidates": 0},
            "fitted_preprocessing_parameters": {},
            "preprocessing": {"impute": "none", "scale": True, "select_k": 0, "scaler": "standard",
                              "selection": "none", "f_threshold": 5, "forward_tol": 0, "pca": 0,
                              "resample": "none"},
            "importance_repeats": 2, "failures": [], "warnings": [],
            "question": "", "language": "zh", "model_backend": "local-registry/v1",
            "model_parameters": {}, "versions": {"scikit-learn": "1.5.0"},
            "inputs": {}, "arguments": {}, "estimator_parameters": {},
        }
        (root / "summary.json").write_text(json.dumps(summary), encoding="utf-8")
        rows = []
        for tag in models:
            for fold, score in enumerate((0.2, 0.25, 0.3), 1):
                rows.append({"model": tag, "fold": fold, "score": score, "train_score": score - 0.05})
        pd.DataFrame(rows).to_csv(root / "cv-scores.csv", index=False)
        pred = pd.DataFrame({
            "sample_id": [f"{i:03}" for i in range(20)],
            "observed": np.linspace(0, 1, 20),
            "predicted": np.linspace(0.05, 1.05, 20),
        })
        pred.to_csv(root / "test-predictions.csv", index=False)
        pd.DataFrame({"sample_id": pred.sample_id, "partition": "test"}).to_csv(
            root / "split-membership.csv", index=False)
        pd.DataFrame({
            "metric": ["mae"], "estimate": [0.25], "lower": [0.1], "upper": [0.4],
            "baseline": [0.5], "gain": [0.25], "gain_lower": [0.1], "gain_upper": [0.4],
        }).to_csv(root / "metric-intervals.csv", index=False)
        analysis = {
            "schema": "psytrainer-diagnostics/v1",
            "interval": {"method": "row percentile bootstrap", "requested_repeats": 100,
                         "units": 20, "omitted_reason": None},
            "findings": [{"code": "baseline", "zh": "基线比较", "en": "baseline", "source": "metric-intervals.csv"}],
            "test_n": 20, "limitations": [],
        }
        (root / "analysis.json").write_text(json.dumps(analysis), encoding="utf-8")
        return summary

    def test_inline_report_has_per_model_sections_and_appendix_table(self):
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            models = ["ModelLRRegressor", "ModelRandomForestRegressor"]
            self.assertLess(len(models), PER_MODEL_SPLIT_THRESHOLD)
            self._write_run(root, models)
            fake_captions = [{"name": "cv-comparison", "caption": "caption", "source": "cv-scores.csv"}]
            figures = root / "figures"
            figures.mkdir()
            # Minimal PNG so the figure loop can open the file.
            from PIL import Image
            Image.new("RGB", (40, 30), color=(255, 255, 255)).save(figures / "cv-comparison.png")
            with patch.object(pipeline_report, "plots", return_value=fake_captions), \
                 patch("pipeline_analysis.analyze", return_value=json.loads((root / "analysis.json").read_text())):
                pipeline_report.generate_report(root)
            text = (root / "report.md").read_text(encoding="utf-8")
            for section in ("## 摘要", "## 要点", "## 1 方法", "### 2.1 模型比较", "### 2.2 测试集表现",
                            "## 3 讨论", "附录 A 各候选模型的详细结果"):
                self.assertIn(section, text)
            self.assertIn("**结论**", text)
            self.assertIn("预测能力", text)
            self.assertIn("DummyRegressor", text)
            self.assertIn("训练集", text)
            for term in ("开发集", "配对改善", "正向置换重要性"):
                self.assertNotIn(term, text)
            self.assertIn("### 线性回归（最终模型）", text)
            self.assertIn("### 随机森林回归器", text)
            for label in ("主张：", "证据：", "阐释：", "lightgbm None"):
                self.assertNotIn(label, text)
            # Markdown tables must be contiguous: no blank line between rows.
            lines = text.splitlines()
            for i, line in enumerate(lines[:-2]):
                if line.startswith("|") and lines[i + 1] == "":
                    self.assertFalse(lines[i + 2].startswith("|"), f"broken table near line {i}")
            self.assertFalse((root / "reports").exists())
            # The Word file is written to its own path and never over a figure.
            self.assertEqual(Image.open(figures / "cv-comparison.png").format, "PNG")
            from docx import Document
            doc = Document(root / "report.docx")
            headings = [p.text for p in doc.paragraphs if p.style.name.startswith("Heading")]
            self.assertIn("2.1 模型比较", headings)
            self.assertTrue(any("线性回归" in h for h in headings))
            self.assertGreaterEqual(len(doc.tables), 3)

    def test_context_and_language_override_on_regeneration(self):
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            self._write_run(root, ["ModelLRRegressor", "ModelRandomForestRegressor"])
            context = root / "ctx.json"
            context.write_text(json.dumps({"target_label": "Depression score", "target_unit": "points",
                                           "tolerance": 0.3}), encoding="utf-8")
            with patch.object(pipeline_report, "plots", return_value=[]), \
                 patch("pipeline_analysis.analyze", return_value=json.loads((root / "analysis.json").read_text())):
                pipeline_report.generate_report(root, context=context, language="en")
                text = (root / "report.md").read_text(encoding="utf-8")
                self.assertIn("Depression score", text)
                self.assertIn("## Abstract", text)
                self.assertIn("had absolute error ≤ 0.30 points", text)
                self.assertTrue((root / "report-context.json").is_file())
                pipeline_report.generate_report(root)  # saved context persists
            self.assertIn("Depression score", (root / "report.md").read_text(encoding="utf-8"))
            bad = root / "bad.json"
            bad.write_text(json.dumps({"colour": "blue"}), encoding="utf-8")
            with self.assertRaises(ValueError):
                pipeline_report.generate_report(root, context=bad)

    def test_large_runs_write_per_model_reports_and_index(self):
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            models = ["ModelLRRegressor", "ModelRandomForestRegressor",
                      "ModelKNeighborsRegressor", "ModelSVRRegressor"]
            self._write_run(root, models)
            fake_captions = []
            with patch.object(pipeline_report, "plots", return_value=fake_captions), \
                 patch("pipeline_analysis.analyze", return_value=json.loads((root / "analysis.json").read_text())):
                pipeline_report.generate_report(root)
            reports = root / "reports"
            self.assertTrue(reports.is_dir())
            for tag in models:
                self.assertTrue((reports / f"{tag}.md").is_file())
                self.assertTrue((reports / f"{tag}.docx").is_file())
            overview = (root / "report.md").read_text(encoding="utf-8")
            self.assertIn("reports/ModelLRRegressor.md", overview)
            self.assertIn("单独成文", overview)
            detail = (reports / "ModelRandomForestRegressor.md").read_text(encoding="utf-8")
            self.assertIn("逐折结果", detail)
            self.assertIn("交叉验证", detail)


if __name__ == "__main__":
    unittest.main()
