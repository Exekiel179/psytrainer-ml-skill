#!/usr/bin/env python3
"""Deterministic figures and Word/Markdown reports from measured Pipeline outputs."""

from __future__ import annotations

import json
import textwrap
from pathlib import Path

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from matplotlib import font_manager
import numpy as np
import pandas as pd
from sklearn.metrics import confusion_matrix, precision_recall_curve, roc_curve

BLUE, ORANGE, GRAY = "#0072B2", "#D55E00", "#656565"


def style():
    installed = {f.name for f in font_manager.fontManager.ttflist}
    fonts = [f for f in ("Arial", "Noto Sans CJK SC", "Microsoft YaHei", "PingFang SC", "DejaVu Sans")
             if f in installed]
    plt.rcParams.update({"font.family": "sans-serif", "font.sans-serif": fonts,
                         "font.size": 9, "axes.labelsize": 9, "xtick.labelsize": 8,
                         "ytick.labelsize": 8, "legend.fontsize": 8,
                         "axes.spines.top": False, "axes.spines.right": False,
                         "axes.linewidth": 0.8, "legend.frameon": False,
                         "svg.fonttype": "none", "pdf.fonttype": 42,
                         "figure.facecolor": "white", "axes.facecolor": "white"})


def save_figure(fig, directory, name):
    fig.canvas.draw()
    # Optional local publication QA; ordinary installs need no external skills.
    import os
    import sys
    if os.environ.get("PSYTRAINER_FIGURE_QA_PATH"):
        sys.path.insert(0, os.environ["PSYTRAINER_FIGURE_QA_PATH"])
        from audit_panel_alignment import require_matplotlib_panel_alignment
        require_matplotlib_panel_alignment(fig, json_out=str(directory / f"{name}.alignment.json"),
                                           tolerance_pt=1.5, gutter_tolerance_pt=1.5, strict=True)
    fig.savefig(directory / f"{name}.png", dpi=600)
    fig.savefig(directory / f"{name}.svg")
    fig.savefig(directory / f"{name}.pdf")
    plt.close(fig)


def plots(directory, summary):
    style()
    figures = directory / "figures"
    figures.mkdir(exist_ok=True)
    cv = pd.read_csv(directory / "cv-scores.csv")
    pred = pd.read_csv(directory / "test-predictions.csv", dtype={"sample_id": str, "observed": str, "predicted": str})
    zh = summary["language"] == "zh"
    n, folds = summary["test_n"], summary["cv_folds"]
    captions = []

    def add(fig, name, caption, source):
        save_figure(fig, figures, name)
        captions.append({"name": name, "caption": caption, "source": source})

    if summary.get("search", {}).get("method", "none") != "none":
        trials = json.loads((directory / "search-results.json").read_text())
        # Paginate models to keep labels readable; all candidate values stay in JSON.
        names = list(dict.fromkeys(row["model"] for row in trials))
        for page in range(0, len(names), 3):
            fig, axes = plt.subplots(len(names[page:page + 3]), 1,
                figsize=(7.09, 2.0 * len(names[page:page + 3])), squeeze=False, layout="constrained")
            for ax, name in zip(axes.ravel(), names[page:page + 3]):
                valid = [row for row in trials if row["model"] == name and row["status"] == "completed"]
                ax.scatter([row["candidate"] for row in valid], [row["mean"] for row in valid], color=BLUE, s=20)
                failed = sum(row["model"] == name and row["status"] == "failed" for row in trials)
                ax.set(title=f"{name.removeprefix('Model')} (failed: {failed})",
                       xlabel="Candidate index", ylabel=summary["metric"])
            add(fig, f"parameter-search-{page // 3 + 1}",
                ("参数搜索：每点为一个候选参数配置的平均验证得分，候选编号对应源文件。搜索使用开发集，分数用于选择而非无偏性能估计；失败数单独列出。" if zh else
                 "Parameter search: each point is mean development CV score for a candidate indexed in the source file. Scores are for selection, not unbiased performance estimates; failed candidates are counted separately."),
                "search-results.json")

    if (directory / "baseline-cv.csv").exists():
        from pipeline_diagnostics_plots import diagnostic_plots
        diagnostic_plots(directory, summary, add)
        (figures / "figure-manifest.json").write_text(json.dumps({"width_mm": 180, "dpi": 600,
            "formats": ["png", "svg", "pdf"], "style": "Nature-inspired; no journal compliance certification",
            "figures": captions}, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
        return captions

    fig, ax = plt.subplots(figsize=(7.09, 3.6), layout="constrained")
    for i, row in enumerate(summary["comparison"]):
        scores = cv.loc[cv.model == row["model"], "score"].to_numpy()
        ax.scatter(scores, np.full(len(scores), i) + np.linspace(-0.1, 0.1, len(scores)),
                   color=GRAY, alpha=0.6, s=18)
        ax.errorbar(row["mean"], i, xerr=row["sd"], fmt="o", color=BLUE, capsize=3)
    ax.set_yticks(range(len(summary["comparison"])),
                  [textwrap.fill(r["model"].removeprefix("Model"), 24) for r in summary["comparison"]])
    ax.invert_yaxis()
    ax.set_ylim(len(summary["comparison"]) - 0.5, -0.5)
    ax.set_xlabel(f"{summary['metric']} ({summary['direction']} is better)")
    add(fig, "cv-comparison",
        (f"模型选择：{folds} 折验证得分。灰点为各折，蓝点和误差线为均值及折间样本标准差（非置信区间）。各模型使用相同划分。"
         if zh else f"Model selection across {folds} folds. Gray points: individual folds; blue: mean and sample SD, not a confidence interval. Splits are shared across models."),
        "cv-scores.csv")

    if summary["task"] == "regression":
        observed = pd.to_numeric(pred.observed)
        predicted = pd.to_numeric(pred.predicted)
        fig, ax = plt.subplots(figsize=(7.09, 4.0), layout="constrained")
        ax.scatter(observed, predicted, color=BLUE, s=16, alpha=0.65, rasterized=True)
        lower, upper = min(observed.min(), predicted.min()), max(observed.max(), predicted.max())
        ax.plot([lower, upper], [lower, upper], color=GRAY, linestyle="--", linewidth=1)
        ax.set(xlabel="Observed outcome", ylabel="Predicted outcome")
        add(fig, "observed-predicted",
            (f"独立测试集预测表现（n={n} 个样本）。每点为一个样本，虚线为理想预测线；未绘制推断区间。"
             if zh else f"Held-out predictions (n={n} samples). Each point is one sample; dashed line indicates ideal predictions. No inferential intervals."),
            "test-predictions.csv")
        fig, ax = plt.subplots(figsize=(7.09, 4.0), layout="constrained")
        ax.scatter(predicted, observed - predicted, color=ORANGE, s=16, alpha=0.65, rasterized=True)
        ax.axhline(0, color=GRAY, linestyle="--", linewidth=1)
        ax.set(xlabel="Predicted outcome", ylabel="Residual (observed - predicted)")
        add(fig, "residuals",
            (f"独立测试集残差诊断（n={n}）。零线表示无预测误差；散点用于检查系统偏差和误差随预测值变化的模式。"
             if zh else f"Held-out residual diagnostics (n={n}). Zero denotes no prediction error. Points show bias and variation across predicted outcomes."),
            "test-predictions.csv")
    else:
        classes = summary["classes"]
        matrix = confusion_matrix(pred.observed, pred.predicted, labels=classes)
        pd.DataFrame(matrix, index=classes, columns=classes).to_csv(directory / "confusion-matrix.csv", index_label="observed")
        fig, ax = plt.subplots(figsize=(7.09, 4.5), layout="constrained")
        ax.imshow(matrix, cmap="Blues", aspect="equal")
        # Numeric class keys keep long labels out of a fixed-size matrix.
        ax.set_xticks(range(len(classes)), [f"C{i}" for i in range(len(classes))])
        ax.set_yticks(range(len(classes)), [f"C{i}" for i in range(len(classes))])
        if len(classes) <= 12:
            for (i, j), value in np.ndenumerate(matrix):
                ax.text(j, i, str(value), ha="center", va="center",
                        color="white" if value > matrix.max() / 2 else "black", fontsize=8)
        ax.set(xlabel="Predicted class", ylabel="Observed class")
        mapping = "; ".join(f"C{i}={value}" for i, value in enumerate(classes))
        add(fig, "confusion-matrix",
            ((f"独立测试集混淆矩阵（n={n}），单元格为样本数。类别映射：" if zh else
              f"Held-out confusion matrix (n={n}); cells contain sample counts. Class key: ") + mapping),
            "confusion-matrix.csv; test-predictions.csv")
        if "positive_score" in pred and pred.observed.nunique() == 2:
            binary = pred.observed == classes[1]
            fpr, tpr, _ = roc_curve(binary, pred.positive_score)
            precision, recall, _ = precision_recall_curve(binary, pred.positive_score)
            pd.DataFrame({"false_positive_rate": fpr, "true_positive_rate": tpr}).to_csv(directory / "roc.csv", index=False)
            pd.DataFrame({"recall": recall, "precision": precision}).to_csv(directory / "precision-recall.csv", index=False)
            for name, xx, yy, xlabel, ylabel, baseline in (
                ("roc", fpr, tpr, "False positive rate", "True positive rate", None),
                ("precision-recall", recall, precision, "Recall", "Precision", float(binary.mean())),
            ):
                fig, ax = plt.subplots(figsize=(7.09, 4.0), layout="constrained")
                if name == "roc":
                    ax.plot(xx, yy, color=BLUE, linewidth=1.4)
                else:
                    ax.step(xx, yy, where="post", color=BLUE, linewidth=1.4)
                if baseline is None:
                    ax.plot([0, 1], [0, 1], linestyle="--", color=GRAY, linewidth=1)
                else:
                    ax.axhline(baseline, linestyle="--", color=GRAY, linewidth=1)
                ax.set(xlabel=xlabel, ylabel=ylabel, xlim=(-0.02, 1.02), ylim=(-0.02, 1.02))
                add(fig, name,
                    (f"独立测试集 {name} 曲线（n={n}，阳性类别={classes[1]}）。虚线为随机基线，曲线无置信区间。"
                     if zh else f"Held-out {name} curve (n={n}; positive class={classes[1]}). Dashed line: chance baseline. No confidence interval."),
                    f"{name}.csv; test-predictions.csv")

    if (directory / "importance.csv").exists():
        all_importance = pd.read_csv(directory / "importance.csv", dtype={"feature": str})
        imp = all_importance.head(15).iloc[::-1]
        fig, ax = plt.subplots(figsize=(7.09, max(3.6, 0.3 * len(imp) + 1.2)), layout="constrained")
        ax.errorbar(imp["mean"], np.arange(len(imp)), xerr=imp.sd, fmt="o", color=BLUE, capsize=3)
        ax.set_yticks(range(len(imp)), [textwrap.fill(s, 25) for s in imp.feature])
        ax.axvline(0, color=GRAY, linestyle="--", linewidth=1)
        ax.set_xlabel(f"Decrease in {summary['metric']} scorer (higher = greater reliance)")
        add(fig, "permutation-importance",
            (f"测试集置换重要性（n={n}，{summary['importance_repeats']} 次置换）。展示前 {len(imp)}/{len(all_importance)} 个特征；全部数值保存在源数据。误差线为置换次数间总体标准差，非置信区间。使用越大越好的评分，损失指标取负值。重要性不表示因果效应。"
             if zh else f"Test permutation importance (n={n}; {summary['importance_repeats']} repeats). Top {len(imp)}/{len(all_importance)} features displayed; all values retained in source data. Bars: population SD across permutations, not confidence intervals. Scorers are oriented higher-is-better (losses negated). Importance is not causal."),
            "importance.csv; importance-repeats.csv")
    (figures / "figure-manifest.json").write_text(json.dumps({"width_mm": 180, "dpi": 600,
        "formats": ["png", "svg", "pdf"], "style": "Nature-inspired; no journal compliance certification",
        "figures": captions}, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    return captions


def generate_report(directory):
    from docx import Document
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Inches, Pt, RGBColor
    from report_explanations import (introduction, preprocessing, metric_explanation, figure_guide,
                                     abstract, validation_methods, result_summary, discussion, model_name,
                                     uncertainty_description)

    directory = Path(directory)
    s = json.loads((directory / "summary.json").read_text(encoding="utf-8"))
    from pipeline_analysis import analyze
    analysis = analyze(directory, s)
    captions = plots(directory, s)
    zh = s["language"] == "zh"
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Inches(8.27), Inches(11.69)
    section.top_margin = section.bottom_margin = Inches(0.7)
    section.left_margin = section.right_margin = Inches(0.7)
    installed = {f.name for f in font_manager.fontManager.ttflist}
    cjk_font = next((f for f in ("Noto Sans CJK SC", "Microsoft YaHei", "Arial Unicode MS", "PingFang SC", "Heiti SC")
                     if f in installed), "Microsoft YaHei")
    for name in ("Normal", "Title", "Heading 1", "Heading 2", "Caption"):
        st = doc.styles[name]
        st.font.name = "Arial"
        fonts = st.element.get_or_add_rPr().rFonts
        for key in list(fonts.attrib):
            if "theme" in key.lower():
                del fonts.attrib[key]
        fonts.set(qn("w:eastAsia"), cjk_font)
        for border in st.element.xpath("./w:pPr/w:pBdr"):
            border.getparent().remove(border)
        st.font.color.rgb = RGBColor.from_string("202124")
    normal = doc.styles["Normal"]
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.15
    doc.styles["Title"].font.size = Pt(22)
    doc.styles["Heading 1"].font.size = Pt(14)
    doc.styles["Caption"].font.size = Pt(9)
    title = f"{s['target']} 预测模型分析报告" if zh else f"Predictive analysis of {s['target']}"
    doc.core_properties.title = title
    doc.core_properties.author = "PsyTrainer ML"
    md = []

    def heading(text, level=1, new_page=False):
        item = doc.add_heading(text, level)
        item.paragraph_format.page_break_before = new_page
        md.append(f"{'#' * (level + 1)} {text}\n")

    def paragraph(text):
        doc.add_paragraph(text)
        md.append(text + "\n")

    doc.add_heading(title, 0)
    md.append(f"# {title}\n")
    heading("摘要" if zh else "Abstract")
    if s["question"]:
        paragraph(s["question"])
    paragraph(abstract(s, zh))
    heading("方法" if zh else "Methods")
    heading("研究目标与数据" if zh else "Objective and Data", 2)
    for text in introduction(s, zh):
        paragraph(text)
    heading("数据划分与验证设计" if zh else "Validation Design", 2)
    paragraph(' '.join(validation_methods(s, zh)))
    heading("预处理与特征构建" if zh else "Preprocessing and Feature Construction", 2)
    processing = preprocessing(s, zh)
    paragraph(' '.join(processing[:-1]))
    paragraph(processing[-1])
    heading("候选模型与选择准则" if zh else "Candidate Models and Selection", 2)
    paragraph((f"本次比较 {len(s['comparison'])} 个有可用验证结果的模型：" if zh else
               f"This run compares {len(s['comparison'])} models with usable validation results: ") +
              "; ".join(f"{model_name(row['model'], zh)} ({s['metric']}={row['mean']:.4g})" for row in s['comparison']))
    paragraph("各模型使用相同验证划分，依据主要指标的平均验证得分选择最终模型。模型标识与完整参数保存在运行记录中。" if zh else
              "Candidates share validation splits. The mean validation score on the primary metric selects the final model; model identifiers and full parameters are retained in the run record.")
    search = s.get("search", {"method": "none"})
    if search["method"] == "none":
        paragraph("候选模型使用默认或显式指定的固定参数，不执行参数搜索。" if zh else
                  "Candidates use default or explicitly configured fixed parameters; no parameter search is performed.")
    else:
        paragraph((f"采用{'网格搜索，依次比较预设配置' if search['method'] == 'grid' else '随机搜索，在预算内抽取预设配置'}，共评估 {search['evaluated_candidates']} 个候选配置，其中 {search['failed_candidates']} 个失败。配置中的超参数控制模型学习或预处理方式，其选择依据开发集验证结果。具体设置见 search-results.json 和 summary.json。" if zh else
                   f"The {search['method']} search evaluated {search['evaluated_candidates']} predefined configurations, with {search['failed_candidates']} failures. Hyperparameters control model learning or preprocessing and are selected using development validation. Exact settings are in search-results.json and summary.json."))
    if s["preprocessing"].get("selection") == "forward":
        paragraph("逐步前向筛选仅接收当前训练分区，内部交叉验证采用相同划分设计，并在内部训练折重新拟合预处理及重采样。" if zh else
                  "Forward selection sees only the current training partition; its internal CV follows the same split design and refits preprocessing and resampling within internal training folds.")
    intervals = pd.read_csv(directory / "metric-intervals.csv")
    records = intervals.astype(object).where(pd.notna(intervals), None).to_dict('records')
    heading("结果" if zh else "Results")
    heading("测试集表现与基线比较" if zh else "Test Performance and Baseline Comparison", 2)
    paragraph(result_summary(s, records, zh))
    table = doc.add_table(rows=1, cols=2)
    table.style = "Light Shading Accent 1"
    for cell, text in zip(table.rows[0].cells, ("指标" if zh else "Metric", "数值" if zh else "Value")):
        cell.text = text
    for key, value in s["test_metrics"].items():
        row = table.add_row().cells
        row[0].text, row[1].text = key, f"{value:.6g}" if value is not None else "undefined"
    md.append("| Metric | Value |\n|---|---|\n" + "\n".join(
        f"| {k} | {v:.6g} |" if v is not None else f"| {k} | undefined |"
        for k, v in s["test_metrics"].items()) + "\n")
    strategy = s.get('baseline', {}).get('strategy', '')
    baseline = {"mean": ("始终预测开发集目标的平均值", "always predicts the development outcome mean"),
                "prior": ("按开发集类别比例给出概率，并按最多见类别作类别预测", "uses development class frequencies as probabilities and predicts the most frequent class"),
                "most_frequent": ("始终预测开发集中最多见的类别", "always predicts the most frequent development class")}.get(strategy)
    paragraph((("简单基线" + baseline[0] + "。" if baseline else "简单基线只使用开发数据学习简单预测规则。") +
               "该参照不利用预测变量，用于量化模型相对于简单预测规则的增益。" if zh else
               ("The dummy baseline " + baseline[1] + ". " if baseline else "The dummy learns a simple rule from development data. ") +
               "It does not use predictors and quantifies the model's gain over a simple prediction rule."))

    heading("估计不确定性" if zh else "Estimate Uncertainty", 2)
    interval = analysis["interval"]
    for text in uncertainty_description(interval, zh):
        paragraph(text)
    table = doc.add_table(rows=1, cols=4)
    table.style = "Light Shading Accent 1"
    headers = ["指标", "模型点估计", "模型 95% 区间", "相对基线增益"] if zh else ["Metric", "Estimate", "Model 95% interval", "Gain over baseline"]
    for cell, text in zip(table.rows[0].cells, headers):
        cell.text = text
    md.append("| " + " | ".join(headers) + " |\n|---|---|---|---|\n")
    md_rows = []
    for row in intervals.itertuples():
        values = [row.metric, f"{row.estimate:.4g}" if pd.notna(row.estimate) else "NA",
                  f"[{row.lower:.4g}, {row.upper:.4g}]" if pd.notna(row.lower) else "NA",
                  f"{row.gain:.4g}" if pd.notna(row.gain) else "NA"]
        for cell, text in zip(table.add_row().cells, values):
            cell.text = text
        md_rows.append("| " + " | ".join(values) + " |")
    md[-1] += "\n".join(md_rows) + "\n"
    paragraph("差值按改善方向定义：得分型指标为模型减基线，误差型指标为基线减模型，因此正值均表示改善。" if zh else
              "Differences are oriented toward improvement: model minus baseline for scores, baseline minus model for losses. Positive values therefore favor the model.")

    for number, figure in enumerate(captions, 1):
        figure_title, explanations = figure_guide(figure, s, analysis, zh)
        heading((f"图 {number} " if zh else f"Figure {number} ") + figure_title, new_page=True)
        path = directory / "figures" / f"{figure['name']}.png"
        from PIL import Image
        with Image.open(path) as image:
            width, height = image.size
        display_width = min(6.7, 4.5 * width / height)
        doc.add_picture(str(path), width=Inches(display_width))
        doc.add_paragraph(figure["caption"], style="Caption")
        paragraph(("源数据：" if zh else "Source data: ") + figure["source"])
        md.append(f"![{figure['name']}](figures/{figure['name']}.png)\n\n{figure['caption']}\n")
        for _, text in explanations:
            paragraph(text)

    heading("讨论与局限" if zh else "Discussion and Limitations", new_page=True)
    for text in discussion(s, analysis, zh):
        paragraph(text)
    if s["failures"]:
        paragraph(("未完成的候选模型：" if zh else "Failed candidates: ") + json.dumps(s["failures"], ensure_ascii=False))
    for warning in s["warnings"]:
        if "omitted" in warning or "every development class" in warning:
            if "Pearson correlation is undefined for a constant dummy" not in warning:
                paragraph(warning)
    heading("附录 评价指标释义" if zh else "Appendix Metric Definitions", new_page=True)
    paragraph("以下说明给出指标定义及其在本次测试集中的数值。NA 或 undefined 表示数据或设计不足以计算该指标。" if zh else
              "The definitions below accompany values from this test set. NA or undefined denotes an estimate unsupported by the data or design.")
    for key, value in s['test_metrics'].items():
        label, text = metric_explanation(key, value, zh)
        heading(label, 2)
        paragraph(text)
    heading("复现记录" if zh else "Reproducibility Record")
    paragraph(("完整配置、输入 SHA256 和软件版本位于 summary.json；划分清单位于 split-membership.csv 和 cv-membership.csv。保存的 pipeline.joblib 包含最终预处理与模型，预测时整体复用。绘图源数据、矢量图与高清预览随报告保存。"
               if zh else "summary.json contains configuration, input SHA256 hashes and software versions. split-membership.csv and cv-membership.csv identify all splits. pipeline.joblib stores the complete fitted preprocessing and model. Figure source data, vectors and high-resolution previews accompany this report."))
    paragraph("; ".join(f"{k} {v}" for k, v in s["versions"].items()))
    # Keep header rows and short data rows intact across Word pagination.
    for table in doc.tables:
        repeat = OxmlElement("w:tblHeader")
        table.rows[0]._tr.get_or_add_trPr().append(repeat)
        for row in table.rows:
            row._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
        if len(table.rows) <= 12:
            for row in table.rows[:-1]:
                for cell in row.cells:
                    for paragraph in cell.paragraphs:
                        paragraph.paragraph_format.keep_with_next = True
    # Direct formatting prevents table/theme fonts from overriding CJK glyphs.
    paragraphs = list(doc.paragraphs) + [p for table in doc.tables for row in table.rows
                                       for cell in row.cells for p in cell.paragraphs]
    for p in paragraphs:
        for run in p.runs:
            run.font.name = "Arial"
            fonts = run._element.get_or_add_rPr().rFonts
            fonts.set(qn("w:eastAsia"), cjk_font)
    doc.save(directory / "report.docx")
    (directory / "report.md").write_text("\n".join(md), encoding="utf-8")
