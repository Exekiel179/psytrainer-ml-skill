"""Tiny block document rendered identically to Markdown and Word.

Keeping one block list guarantees the .md and .docx reports carry the same
content, numbering and tables.
"""

from __future__ import annotations

import re
from pathlib import Path

from report_narrative import tidy

BOLD = re.compile(r"\*\*(.+?)\*\*")


class ReportDocument:
    def __init__(self, title, zh):
        self.title, self.zh = tidy(title, zh), zh
        self.blocks = []
        self.figure_count = 0
        self.table_count = 0

    def heading(self, text, level=1, page_break=False):
        self.blocks.append(("heading", level, text, page_break))

    def para(self, text, style=None):
        if text:
            self.blocks.append(("para", tidy(text, self.zh), style))

    def labelled(self, label, text):
        """Paragraph with a bold run-in label, e.g. structured abstracts."""
        if text:
            self.blocks.append(("para", f"**{label}**{'：' if self.zh else ': '}{tidy(text, self.zh)}", None))

    def bullets(self, items, ordered=False):
        items = [tidy(i, self.zh) for i in items if i]
        if items:
            self.blocks.append(("bullets", items, ordered))

    def table(self, caption, headers, rows, note=None):
        self.table_count += 1
        label = f"表 {self.table_count}" if self.zh else f"Table {self.table_count}"
        if note:
            note = tidy(("注：" if self.zh else "Note: ") + note, self.zh)
        self.blocks.append(("table", f"{label} {tidy(caption, self.zh)}" if self.zh else f"{label}. {caption}",
                            [tidy(h, self.zh) for h in headers], [[tidy(str(c), self.zh) for c in r] for r in rows], note))
        return label

    def figure(self, path, caption, source):
        self.figure_count += 1
        label = f"图 {self.figure_count}" if self.zh else f"Figure {self.figure_count}"
        source_text = ("（数据：" if self.zh else "(Data: ") + source + ("）" if self.zh else ")")
        self.blocks.append(("figure", Path(path), label, tidy(caption, self.zh), source_text))
        return label

    def next_figure_label(self, offset=1):
        n = self.figure_count + offset
        return f"图 {n}" if self.zh else f"Figure {n}"

    # ------------------------------------------------------------ markdown
    def markdown(self, directory):
        out = [f"# {self.title}", ""]
        for block in self.blocks:
            kind = block[0]
            if kind == "heading":
                out += [f"{'#' * (block[1] + 1)} {block[2]}", ""]
            elif kind == "para":
                out += [block[1], ""]
            elif kind == "bullets":
                out += [(f"{i}. " if block[2] else "- ") + item for i, item in enumerate(block[1], 1)] + [""]
            elif kind == "table":
                _, caption, headers, rows, note = block
                out += [f"**{caption}**", "", "| " + " | ".join(headers) + " |",
                        "|" + "|".join("---" for _ in headers) + "|"]
                out += ["| " + " | ".join(c.replace("|", "\\|") for c in r) + " |" for r in rows]
                out += [""] + ([note, ""] if note else [])
            elif kind == "figure":
                _, path, label, caption, source = block
                rel = path.relative_to(directory).as_posix() if path.is_absolute() else path.as_posix()
                out += [f"![{label}]({rel})", "", f"**{label}** {caption} {source}", ""]
        return "\n".join(out).rstrip() + "\n"

    # ------------------------------------------------------------ word
    def docx(self, path, cjk_font):
        from docx import Document
        from docx.enum.table import WD_TABLE_ALIGNMENT
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from docx.oxml import OxmlElement
        from docx.oxml.ns import qn
        from docx.shared import Inches, Pt, RGBColor
        from PIL import Image

        doc = Document()
        section = doc.sections[0]
        section.page_width, section.page_height = Inches(8.27), Inches(11.69)
        section.top_margin = section.bottom_margin = Inches(0.8)
        section.left_margin = section.right_margin = Inches(0.9)
        latin = "Times New Roman" if self.zh else "Calibri"
        for name in ("Normal", "Title", "Heading 1", "Heading 2", "Heading 3", "Caption", "List Bullet", "List Number"):
            style = doc.styles[name]
            style.font.name = latin
            fonts = style.element.get_or_add_rPr().rFonts
            for key in list(fonts.attrib):
                if "theme" in key.lower():
                    del fonts.attrib[key]
            fonts.set(qn("w:eastAsia"), cjk_font)
            for border in style.element.xpath("./w:pPr/w:pBdr"):
                border.getparent().remove(border)
            style.font.color.rgb = RGBColor.from_string("1F1F1F")
        normal = doc.styles["Normal"]
        normal.font.size = Pt(10.5)
        normal.paragraph_format.space_after = Pt(6)
        normal.paragraph_format.line_spacing = 1.3
        doc.styles["Title"].font.size = Pt(18)
        doc.styles["Title"].font.bold = True
        for name, size in (("Heading 1", 14), ("Heading 2", 12), ("Heading 3", 11)):
            doc.styles[name].font.size = Pt(size)
            doc.styles[name].font.bold = True
            doc.styles[name].paragraph_format.space_before = Pt(12)
            doc.styles[name].paragraph_format.space_after = Pt(6)
        doc.styles["Caption"].font.size = Pt(9)
        doc.styles["Caption"].font.italic = False
        doc.styles["Caption"].font.color.rgb = RGBColor.from_string("3C3C3C")
        doc.core_properties.title = self.title
        doc.core_properties.author = "PsyTrainer ML"

        def runs(paragraph, text, size=None):
            for i, part in enumerate(BOLD.split(text)):
                if part:
                    run = paragraph.add_run(part)
                    run.bold = bool(i % 2)
                    if size:
                        run.font.size = Pt(size)

        def three_line(table):
            # Academic three-line table: thick top/bottom rules, thin header rule.
            borders = OxmlElement("w:tblBorders")
            for edge, size in (("top", 12), ("bottom", 12)):
                element = OxmlElement(f"w:{edge}")
                element.set(qn("w:val"), "single")
                element.set(qn("w:sz"), str(size))
                element.set(qn("w:color"), "000000")
                borders.append(element)
            table._tbl.tblPr.append(borders)
            for cell in table.rows[0].cells:
                tc = cell._tc.get_or_add_tcPr()
                cell_borders = OxmlElement("w:tcBorders")
                bottom = OxmlElement("w:bottom")
                bottom.set(qn("w:val"), "single")
                bottom.set(qn("w:sz"), "6")
                bottom.set(qn("w:color"), "000000")
                cell_borders.append(bottom)
                tc.append(cell_borders)

        title = doc.add_paragraph(style="Title")
        title.alignment = WD_ALIGN_PARAGRAPH.CENTER
        runs(title, self.title)
        for block in self.blocks:
            kind = block[0]
            if kind == "heading":
                item = doc.add_heading(block[2], min(block[1], 3))
                item.paragraph_format.page_break_before = block[3]
                item.paragraph_format.keep_with_next = True
            elif kind == "para":
                paragraph = doc.add_paragraph(style=block[2]) if block[2] else doc.add_paragraph()
                runs(paragraph, block[1])
            elif kind == "bullets":
                for item in block[1]:
                    paragraph = doc.add_paragraph(style="List Number" if block[2] else "List Bullet")
                    runs(paragraph, item)
            elif kind == "table":
                _, caption, headers, rows, note = block
                cap = doc.add_paragraph(style="Caption")
                cap.paragraph_format.keep_with_next = True
                runs(cap, f"**{caption}**")
                table = doc.add_table(rows=1, cols=len(headers))
                table.alignment = WD_TABLE_ALIGNMENT.CENTER
                for cell, text in zip(table.rows[0].cells, headers):
                    cell.text = ""
                    runs(cell.paragraphs[0], f"**{text}**", 9)
                for values in rows:
                    cells = table.add_row().cells
                    for cell, text in zip(cells, values):
                        cell.text = ""
                        runs(cell.paragraphs[0], text, 9)
                three_line(table)
                header = OxmlElement("w:tblHeader")
                table.rows[0]._tr.get_or_add_trPr().append(header)
                for row in table.rows:
                    row._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
                    for cell in row.cells:
                        for p in cell.paragraphs:
                            p.paragraph_format.space_after = Pt(1)
                            p.paragraph_format.line_spacing = 1.0
                            if len(table.rows) <= 15:
                                p.paragraph_format.keep_with_next = True
                if note:
                    runs(doc.add_paragraph(style="Caption"), note)
                else:
                    doc.add_paragraph().paragraph_format.space_after = Pt(2)
            elif kind == "figure":
                _, image_path, label, caption, source = block
                with Image.open(image_path) as image:
                    width, height = image.size
                picture = doc.add_paragraph()
                picture.alignment = WD_ALIGN_PARAGRAPH.CENTER
                picture.paragraph_format.keep_with_next = True
                picture.add_run().add_picture(str(image_path), width=Inches(min(6.4, 5.0 * width / height)))
                runs(doc.add_paragraph(style="Caption"), f"**{label}** {caption} {source}")
        for paragraph in list(doc.paragraphs) + [p for t in doc.tables for r in t.rows for c in r.cells for p in c.paragraphs]:
            for run in paragraph.runs:
                run.font.name = latin
                run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), cjk_font)
        doc.save(path)
